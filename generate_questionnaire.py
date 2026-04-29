from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Add title
title = doc.add_heading('NexClinic Doctor User Feedback Questionnaire', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add header info
header_table = doc.add_table(rows=3, cols=2)
header_table.autofit = False
header_table.allow_autofit = False

header_table.rows[0].cells[0].text = 'Test Date:'
header_table.rows[0].cells[1].text = '________________'
header_table.rows[1].cells[0].text = 'Doctor Name (Optional):'
header_table.rows[1].cells[1].text = '________________'
header_table.rows[2].cells[0].text = 'Experience Level:'
header_table.rows[2].cells[1].text = '☐ Beginner    ☐ Intermediate    ☐ Advanced'

doc.add_paragraph()

# Section 1
doc.add_heading('Section 1: Overall Experience', level=1)

questions_s1 = [
    {
        'num': 1,
        'text': 'How would you rate your overall experience with NexClinic?',
        'options': '☐ Poor (1)  ☐ Fair (2)  ☐ Good (3)  ☐ Very Good (4)  ☐ Excellent (5)'
    },
    {
        'num': 2,
        'text': 'How likely are you to recommend NexClinic to other doctors?',
        'options': '☐ Very Unlikely  ☐ Unlikely  ☐ Neutral  ☐ Likely  ☐ Very Likely'
    },
    {
        'num': 3,
        'text': 'What was your first impression of the system?',
        'options': '_________________________________________________________________'
    }
]

for q in questions_s1:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 2
doc.add_heading('Section 2: Appointment Management', level=1)

questions_s2 = [
    {
        'num': 4,
        'text': 'How easy was it to navigate appointment scheduling?',
        'options': '☐ Very Difficult  ☐ Difficult  ☐ Neutral  ☐ Easy  ☐ Very Easy'
    },
    {
        'num': 5,
        'text': 'Did you successfully complete an appointment booking?',
        'options': '☐ Yes  ☐ No\nIf No, what prevented you? _____________________________________________'
    },
    {
        'num': 6,
        'text': 'How clear were the available appointment slots?',
        'options': '☐ Very Unclear  ☐ Unclear  ☐ Neutral  ☐ Clear  ☐ Very Clear'
    },
    {
        'num': 7,
        'text': 'Could you easily set your availability/schedule?',
        'options': '☐ Yes  ☐ Partially  ☐ No\nFeedback: ___________________________________________________________'
    }
]

for q in questions_s2:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 3
doc.add_heading('Section 3: Patient Management', level=1)

questions_s3 = [
    {
        'num': 8,
        'text': 'How intuitive was viewing your patient list?',
        'options': '☐ Very Difficult  ☐ Difficult  ☐ Neutral  ☐ Easy  ☐ Very Easy'
    },
    {
        'num': 9,
        'text': 'Were you able to view patient appointment history?',
        'options': '☐ Yes  ☐ Partially  ☐ No'
    },
    {
        'num': 10,
        'text': 'How would you rate the organization of patient information?',
        'options': '☐ Poor  ☐ Fair  ☐ Good  ☐ Very Good  ☐ Excellent'
    }
]

for q in questions_s3:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 4
doc.add_heading('Section 4: User Interface & Design', level=1)

questions_s4 = [
    {
        'num': 11,
        'text': 'How would you rate the visual design and layout?',
        'options': '☐ Poor  ☐ Fair  ☐ Good  ☐ Very Good  ☐ Excellent'
    },
    {
        'num': 12,
        'text': 'Was the navigation menu clear and logical?',
        'options': '☐ Very Confusing  ☐ Confusing  ☐ Neutral  ☐ Clear  ☐ Very Clear'
    },
    {
        'num': 13,
        'text': 'Were buttons, icons, and controls easy to understand?',
        'options': '☐ Very Difficult  ☐ Difficult  ☐ Neutral  ☐ Easy  ☐ Very Easy'
    },
    {
        'num': 14,
        'text': 'Did you encounter any confusing features or unclear labels?',
        'options': '☐ Yes  ☐ No\nIf Yes, which ones? __________________________________________________'
    }
]

for q in questions_s4:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 5
doc.add_heading('Section 5: Performance & Technical', level=1)

questions_s5 = [
    {
        'num': 15,
        'text': 'How was the system\'s responsiveness/loading speed?',
        'options': '☐ Very Slow  ☐ Slow  ☐ Average  ☐ Fast  ☐ Very Fast'
    },
    {
        'num': 16,
        'text': 'Did you experience any crashes, errors, or technical issues?',
        'options': '☐ Yes  ☐ No\nIf Yes, describe: ____________________________________________________'
    },
    {
        'num': 17,
        'text': 'How well did the system work on your device?',
        'options': '☐ Poor  ☐ Fair  ☐ Good  ☐ Very Good  ☐ Excellent'
    }
]

for q in questions_s5:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 6
doc.add_heading('Section 6: Features & Functionality', level=1)

questions_s6 = [
    {
        'num': 18,
        'text': 'Which features were most valuable to you?',
        'options': '☐ Appointment Scheduling  ☐ Patient Management  ☐ Availability Setting\n☐ Profile Management  ☐ Other: _______________________________________'
    },
    {
        'num': 19,
        'text': 'What features were missing or would you like to see?',
        'options': '_________________________________________________________________\n_________________________________________________________________'
    },
    {
        'num': 20,
        'text': 'Did all features work as expected?',
        'options': '☐ Yes  ☐ Mostly  ☐ Some didn\'t  ☐ No\nIssues encountered: __________________________________________________'
    }
]

for q in questions_s6:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 7
doc.add_heading('Section 7: Mobile/Responsiveness', level=1)

questions_s7 = [
    {
        'num': 21,
        'text': 'How well does the system work on mobile devices?',
        'options': '☐ Not Tested  ☐ Poor  ☐ Fair  ☐ Good  ☐ Excellent'
    },
    {
        'num': 22,
        'text': 'Would you use this on mobile to manage appointments?',
        'options': '☐ Yes  ☐ Maybe  ☐ No'
    }
]

for q in questions_s7:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 8
doc.add_heading('Section 8: Support & Documentation', level=1)

questions_s8 = [
    {
        'num': 23,
        'text': 'Was help or guidance readily available when needed?',
        'options': '☐ Yes  ☐ Partially  ☐ No'
    },
    {
        'num': 24,
        'text': 'Would you benefit from tutorials or guides?',
        'options': '☐ Yes  ☐ No\nTopics needed: _____________________________________________________'
    }
]

for q in questions_s8:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Section 9
doc.add_heading('Section 9: Open Feedback', level=1)

questions_s9 = [
    {
        'num': 25,
        'text': 'What was the best aspect of using NexClinic?',
        'options': '_________________________________________________________________'
    },
    {
        'num': 26,
        'text': 'What was the most frustrating or challenging aspect?',
        'options': '_________________________________________________________________'
    },
    {
        'num': 27,
        'text': 'What is the single most important improvement we should make?',
        'options': '_________________________________________________________________'
    },
    {
        'num': 28,
        'text': 'Any additional comments or suggestions?',
        'options': '_________________________________________________________________\n_________________________________________________________________'
    }
]

for q in questions_s9:
    p = doc.add_paragraph(f"{q['num']}. {q['text']}", style='List Number')
    doc.add_paragraph(q['options'])

# Add footer
doc.add_paragraph()
footer = doc.add_paragraph('Thank you for your feedback!')
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_format = footer.runs[0].font
footer_format.bold = True

# Save document
output_path = r'd:\UoP\Semester 3\CO2060\cepdnaclk\e23-co2060-NexClinic\NexClinic_Doctor_Feedback_Questionnaire.docx'
doc.save(output_path)
print(f"Document created successfully: {output_path}")
